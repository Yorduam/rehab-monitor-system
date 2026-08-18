# -*- coding: utf-8 -*-
"""Rebuild the six enrollment templates as native Word documents.

The originals were Excel sheets saved as .docx: the whole page was one
borderless 9-10 column grid, no font was declared anywhere (so Word fell back
to the Calibri theme font) and the statement's grid was a third wider than the
printable area. This script emits proper paragraphs, real bordered tables and
Times New Roman, reusing the wording from enrollmentContent.py.

    python scripts/buildEnrollmentTemplates.py
"""
import os
import sys

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Emu, Mm, Pt, RGBColor

sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))
from enrollmentContent import CONTRACT_BODY, PLAN_SERVICES

OUT = os.path.normpath(os.path.join(
    os.path.dirname(os.path.abspath(__file__)), '..', 'templates', 'documents'))

FONT = 'Times New Roman'
BODY_PT = 12
SMALL_PT = 9
INDENT = Cm(1.25)
TEXT_WIDTH = Cm(17)

ORG_FULL = ('Государственное бюджетное учреждение города Москвы '
            '«Центр социальной интеграции Дианы Гурцкая» Департамента труда '
            'и социальной защиты населения города Москвы (ГБУ ЦСИ Дианы Гурцкая)')
ORG_REQUISITES = [
    'Юридический адрес: 105082, г. Москва, вн.тер.г. Муниципальный округ '
    'Басманный, ул. Большая Почтовая, д.35, стр.1',
    'Почтовый адрес: 105082, г. Москва, ул. Большая Почтовая, д.35, стр.1',
    'ОГРН 1187746289257',
    'ИНН / КПП 9701103096 / 770101001',
    'Телефон: +7 (499) 261-00-09',
    'Официальный сайт: https://center-diana.ru/',
    'E-mail: gurtskaya.center@social.mos.ru',
]
SIGN_LINE = '___________________ / А.А.Мартыненко /'


def _borders(pr, edges, sz=4, color='000000', container=False):
    if container:
        el = pr
    else:
        el = pr.find(qn('w:tcBorders'))
        if el is None:
            el = OxmlElement('w:tcBorders')
            pr.append(el)
    for name in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        if name not in edges:
            continue
        edge = el.find(qn('w:' + name))
        if edge is None:
            edge = OxmlElement('w:' + name)
            el.append(edge)
        style = edges[name]
        edge.set(qn('w:val'), 'nil' if style is None else 'single')
        if style is not None:
            edge.set(qn('w:sz'), str(sz if style is True else style))
            edge.set(qn('w:space'), '0')
            edge.set(qn('w:color'), color)


TBL_PR_SEQ = ('w:tblStyle', 'w:tblpPr', 'w:tblOverlap', 'w:bidiVisual',
              'w:tblStyleRowBandSize', 'w:tblStyleColBandSize', 'w:tblW', 'w:jc',
              'w:tblCellSpacing', 'w:tblInd', 'w:tblBorders', 'w:shd',
              'w:tblLayout', 'w:tblCellMar', 'w:tblLook', 'w:tblCaption',
              'w:tblDescription')


def tbl_pr_add(table, tag):
    """Word rejects tblPr children written out of schema order."""
    tblPr = table._tbl.tblPr
    el = tblPr.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        tblPr.insert_element_before(el, *TBL_PR_SEQ[TBL_PR_SEQ.index(tag) + 1:])
    return el


def table_borders(table, inside_h=True):
    edges = {k: True for k in ('top', 'left', 'bottom', 'right', 'insideV')}
    edges['insideH'] = True if inside_h else None
    _borders(tbl_pr_add(table, 'w:tblBorders'), edges, container=True)


TC_PR_SEQ = ('w:cnfStyle', 'w:tcW', 'w:gridSpan', 'w:hMerge', 'w:vMerge',
             'w:tcBorders', 'w:shd', 'w:noWrap', 'w:tcMar', 'w:textDirection',
             'w:tcFitText', 'w:vAlign', 'w:hideMark')


def tc_pr_add(cell, tag):
    tcPr = cell._tc.get_or_add_tcPr()
    el = tcPr.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        tcPr.insert_element_before(el, *TC_PR_SEQ[TC_PR_SEQ.index(tag) + 1:])
    return el


def underline_cell(cell):
    _borders(tc_pr_add(cell, 'w:tcBorders'), {'bottom': True}, container=True)


def shade(cell, fill='F2F2F2'):
    el = tc_pr_add(cell, 'w:shd')
    el.set(qn('w:val'), 'clear')
    el.set(qn('w:fill'), fill)


def cell_margins(table, left=108, right=108, top=40, bottom=40):
    mar = tbl_pr_add(table, 'w:tblCellMar')
    for name, val in (('top', top), ('left', left), ('bottom', bottom), ('right', right)):
        el = OxmlElement('w:' + name)
        el.set(qn('w:w'), str(val))
        el.set(qn('w:type'), 'dxa')
        mar.append(el)


def fixed_layout(table):
    tbl_pr_add(table, 'w:tblLayout').set(qn('w:type'), 'fixed')


P_PR_SEQ = ('w:pStyle', 'w:keepNext', 'w:keepLines', 'w:pageBreakBefore',
            'w:framePr', 'w:widowControl', 'w:numPr', 'w:suppressLineNumbers',
            'w:pBdr', 'w:shd', 'w:tabs', 'w:suppressAutoHyphens', 'w:kinsoku',
            'w:wordWrap', 'w:overflowPunct', 'w:topLinePunct', 'w:autoSpaceDE',
            'w:autoSpaceDN', 'w:bidi', 'w:adjustRightInd', 'w:snapToGrid',
            'w:spacing', 'w:ind', 'w:contextualSpacing', 'w:mirrorIndents',
            'w:suppressOverlap', 'w:jc', 'w:textDirection', 'w:textAlignment',
            'w:textboxTightWrap', 'w:outlineLvl', 'w:divId', 'w:cnfStyle',
            'w:rPr', 'w:sectPr', 'w:pPrChange')


def p_pr_add(paragraph, tag):
    pPr = paragraph._p.get_or_add_pPr()
    el = pPr.find(qn(tag))
    if el is None:
        el = OxmlElement(tag)
        pPr.insert_element_before(el, *P_PR_SEQ[P_PR_SEQ.index(tag) + 1:])
    return el


def keep_together(paragraph):
    p_pr_add(paragraph, 'w:keepNext')
    p_pr_add(paragraph, 'w:keepLines')


def new_document():
    doc = Document()
    normal = doc.styles['Normal']
    normal.font.name = FONT
    normal.font.size = Pt(BODY_PT)
    rpr = normal.element.get_or_add_rPr()
    fonts = rpr.get_or_add_rFonts()
    for attr in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'):
        fonts.set(qn(attr), FONT)
    for attr in ('w:asciiTheme', 'w:hAnsiTheme', 'w:cstheme', 'w:eastAsiaTheme'):
        if fonts.get(qn(attr)) is not None:
            del fonts.attrib[qn(attr)]
    pf = normal.paragraph_format
    pf.space_before = Pt(0)
    pf.space_after = Pt(0)
    pf.line_spacing = 1.0

    section = doc.sections[0]
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(2)
    section.bottom_margin = Cm(2)
    section.left_margin = Cm(2.5)
    section.right_margin = Cm(1.5)
    return doc


def para(host, text='', align=None, bold=False, italic=False, size=BODY_PT,
         first_line=None, left=None, space_before=0, space_after=0, caps=False):
    p = host.add_paragraph()
    if text:
        run = p.add_run(text)
        run.bold = bold
        run.italic = italic
        run.font.size = Pt(size)
        run.font.name = FONT
        if caps:
            run.font.all_caps = True
    pf = p.paragraph_format
    pf.space_before = Pt(space_before)
    pf.space_after = Pt(space_after)
    pf.line_spacing = 1.0
    if align is not None:
        p.alignment = align
    if first_line is not None:
        pf.first_line_indent = first_line
    if left is not None:
        pf.left_indent = left
    return p


def heading(host, text, size=BODY_PT, space_before=10, space_after=6):
    p = para(host, text, align=WD_ALIGN_PARAGRAPH.CENTER, bold=True, size=size,
             space_before=space_before, space_after=space_after)
    keep_together(p)
    return p


def body(host, text):
    return para(host, text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, first_line=INDENT,
                space_after=2)


def caption(host, text, align=WD_ALIGN_PARAGRAPH.CENTER):
    return para(host, text, align=align, italic=True, size=SMALL_PT, space_after=4)


def clear(cell):
    """A fresh cell already owns one empty paragraph — drop it."""
    p = cell.paragraphs[0]._p
    p.getparent().remove(p)


def grid(doc, widths, rows=0):
    table = doc.add_table(rows=rows, cols=len(widths))
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False
    fixed_layout(table)
    twips = [Emu(int(w)).twips for w in widths]
    for col, tw in zip(table._tbl.tblGrid.findall(qn('w:gridCol')), twips):
        col.set(qn('w:w'), str(tw))
    tblW = tbl_pr_add(table, 'w:tblW')
    tblW.set(qn('w:w'), str(sum(twips)))
    tblW.set(qn('w:type'), 'dxa')
    for row in table.rows:
        for cell, w in zip(row.cells, widths):
            cell.width = w
    table._widths = widths
    return table


def add_row(table):
    row = table.add_row()
    for cell, w in zip(row.cells, table._widths):
        cell.width = w
        clear(cell)
    return row


def span_row(table, row):
    """Collapse a row to one cell. python-docx's merge() needs each cell to own
    a paragraph, but add_row() strips them, so write the gridSpan directly."""
    tcs = row._tr.tc_lst
    for tc in tcs[1:]:
        row._tr.remove(tc)
    cell = row.cells[0]
    cell.width = sum(table._widths, Cm(0))
    tc_pr_add(cell, 'w:gridSpan').set(qn('w:val'), str(len(tcs)))
    return cell


def field_row(table, label, value, note=None):
    """label | value-on-a-rule, with an optional italic note under the value."""
    row = add_row(table)
    para(row.cells[0], label, space_after=2)
    para(row.cells[1], value, space_after=0)
    underline_cell(row.cells[1])
    if note:
        note_row = add_row(table)
        para(note_row.cells[0], '')
        caption(note_row.cells[1], note)
    return row


def signature_block(doc, name_token, who='Заказчика', width=TEXT_WIDTH):
    table = grid(doc, [Cm(7), width - Cm(7)])
    cell_margins(table, left=0, right=108)
    row = add_row(table)
    para(row.cells[0], '', space_after=0)
    underline_cell(row.cells[0])
    para(row.cells[1], name_token, align=WD_ALIGN_PARAGRAPH.CENTER, space_after=0)
    underline_cell(row.cells[1])
    row = add_row(table)
    caption(row.cells[0], '(подпись %s)' % who)
    caption(row.cells[1], '(Ф.И.О. %s)' % who)
    return table


STATEMENT_TITLE = [
    'ЗАЯВЛЕНИЕ',
    'о зачислении на реабилитационный курс социальной интеграции',
    'лиц с ограничениями жизнедеятельности',
]
STATEMENT_SUBJECT = (
    'на реабилитационный курс социальной интеграции лиц с ограничениями '
    'жизнедеятельности в Государственное бюджетное учреждение города Москвы '
    '«Центр социальной интеграции Дианы Гурцкая» Департамента труда и социальной '
    'защиты населения города Москвы (ГБУ ЦСИ Дианы Гурцкая) для получения услуг '
    'по социальной интеграции лиц с ограничениями жизнедеятельности '
    '(далее – услуги) в соответствии с индивидуальным планом предоставления услуг '
    'на срок с «____» ______________ 20___ г. по «____» ______________ 20___ г.')


def build_statement(minor):
    doc = new_document()
    left = Cm(8.5)

    for line in ('Директору ГБУ ЦСИ Дианы Гурцкая', 'Ф.В.Молькову'):
        para(doc, line, left=left)
    para(doc, 'от ${fullNameParent}', left=left, space_before=6)
    para(doc, '(Ф.И.О. родителя / законного представителя несовершеннолетнего)'
         if minor else '(Ф.И.О. заявителя)',
         left=left, italic=True, size=SMALL_PT, space_after=6)

    passport = [
        ('паспорт: серия', '${passportSerialParent}'),
        ('номер', '${passportNumberParent}'),
        ('кем выдан:', '${passportWhoParent}'),
        ('дата выдачи:', '${passportDateParent}'),
    ]
    if minor:
        passport.append(('код подразделения:', '${passportCodeParent}'))
    passport += [
        ('адрес регистрации:', '${passportRegistrationParent}'),
        ('тел.:', '${telephoneParent}'),
    ]
    for label, value in passport:
        para(doc, '%s %s' % (label, value), left=left)

    para(doc, '', space_after=10)
    for i, line in enumerate(STATEMENT_TITLE):
        para(doc, line, align=WD_ALIGN_PARAGRAPH.CENTER, bold=(i == 0),
             space_after=2)
    para(doc, '', space_after=8)

    table = grid(doc, [Cm(7), Cm(10)])
    cell_margins(table, left=0)
    if minor:
        row = add_row(table)
        para(row.cells[0], 'Прошу зачислить', space_after=2)
        para(row.cells[1], '${fullName}', space_after=0)
        underline_cell(row.cells[1])
        row = add_row(table)
        caption(row.cells[0], '(сына, дочь, подопечного)', WD_ALIGN_PARAGRAPH.LEFT)
        caption(row.cells[1], '(Ф.И.О. полностью несовершеннолетнего)')
    else:
        field_row(table, 'Прошу зачислить меня,', '${fullName}',
                  '(Ф.И.О. полностью)')

    field_row(table, 'Дата рождения:', '${birthDate}')
    field_row(table,
              'Свидетельство о рождении / паспорт:' if minor else 'Паспорт:',
              '${passportSerial} ${passportNumber}', '(серия, номер)')
    field_row(table, 'Кем выдан:', '${passportWho}')
    field_row(table, 'Адрес регистрации:', '${passportRegistration}')
    field_row(table, 'Адрес фактического места проживания:', '${address}')
    field_row(table, 'Группа инвалидности:', '${disabledGroup}')

    para(doc, '', space_after=8)
    para(doc, STATEMENT_SUBJECT, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=16)

    signature_block(doc, '${fullNameParent}',
                    'родителя / законного представителя' if minor else 'заявителя')
    para(doc, '', space_after=10)
    para(doc, '«_____» ________________________ 20___ г.', space_after=14)

    ack = [
        'С Положением о порядке зачисления на реабилитационный курс социальной '
        'интеграции лиц с ограничениями жизнедеятельности в ГБУ ЦСИ Дианы '
        'Гурцкая ознакомлен(а)',
        'С Положением о пропускном и внутриобъектовом режиме в ГБУ ЦСИ Дианы '
        'Гурцкая ознакомлен(а)',
        'Подтверждаю своё согласие на обработку персональных данных, фото- '
        'и видеосъёмку, а также обнародование и дальнейшее использование '
        + ('изображения своего несовершеннолетнего ребёнка.'
           if minor else 'своего изображения.'),
    ]
    table = grid(doc, [Cm(12), Cm(5)])
    cell_margins(table, left=0)
    for text in ack:
        row = add_row(table)
        para(row.cells[0], text, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_after=0)
        para(row.cells[1], '', space_after=0)
        underline_cell(row.cells[1])
        row = add_row(table)
        para(row.cells[0], '', space_after=8)
        caption(row.cells[1], '(подпись)')
    return doc


def build_plan(minor):
    doc = new_document()
    age = 'minor' if minor else 'adult'
    left = Cm(8.5)

    for line in ('Приложение № 1',
                 'к договору безвозмездного оказания услуг',
                 'по социальной интеграции лиц',
                 'с ограничениями жизнедеятельности',
                 'от «____» _____________ 20___ г. № ____________'):
        para(doc, line, left=left, size=SMALL_PT + 2)

    para(doc, '', space_after=12)
    heading(doc, 'Индивидуальный план', space_before=0, space_after=2)
    heading(doc, 'предоставления услуг по социальной интеграции',
            space_before=0, space_after=12)

    info = grid(doc, [Cm(6), Cm(11)])
    cell_margins(info, left=0)
    if minor:
        field_row(info, 'Получатель услуг:', '${fullName}')
        field_row(info, 'Заказчик:', '${fullNameParent}')
    else:
        field_row(info, 'Заказчик (получатель услуг):', '${fullName}')
    para(doc, '', space_after=10)

    table = grid(doc, [Cm(13), Cm(4)])
    table_borders(table)
    cell_margins(table)
    for kind, text in PLAN_SERVICES[age]:
        row = add_row(table)
        if kind == 'g':
            para(row.cells[0], text, bold=True, align=WD_ALIGN_PARAGRAPH.CENTER)
            para(row.cells[1], 'Количество мероприятий и услуг', bold=True,
                 align=WD_ALIGN_PARAGRAPH.CENTER, size=SMALL_PT + 1)
            shade(row.cells[0])
            shade(row.cells[1])
            keep_together(row.cells[0].paragraphs[0])
        elif kind == 'sub':
            para(span_row(table, row), text, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                 size=SMALL_PT + 2)
        else:
            para(row.cells[0], text, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
                 size=SMALL_PT + 2)
            para(row.cells[1], '', align=WD_ALIGN_PARAGRAPH.CENTER)

    para(doc, '', space_after=16)
    signs = grid(doc, [Cm(8.5), Cm(8.5)])
    cell_margins(signs, left=0)
    row = add_row(signs)
    para(row.cells[0], 'Исполнитель:', bold=True, space_after=4)
    para(row.cells[1], 'Заказчик:', bold=True, space_after=4)

    row = add_row(signs)
    left_cell, right_cell = row.cells
    para(left_cell, 'ГБУ ЦСИ Дианы Гурцкая')
    para(left_cell, 'Заместитель директора', space_after=16)
    para(left_cell, SIGN_LINE)
    para(right_cell, '${fullNameParent}' if minor else '${fullName}',
         space_after=16)
    para(right_cell, '', space_after=0)
    para(right_cell, '___________________', space_after=0)
    caption(right_cell, '(подпись)', WD_ALIGN_PARAGRAPH.LEFT)
    return doc


REQ_ADULT = [
    ('ФИО (полностью):', '${fullName}'),
    ('Адрес регистрации:', '${passportRegistration}'),
    ('Адрес фактический (почтовый):', '${address}'),
    ('Паспортные данные:', None),
    ('серия:', '${passportSerial}'),
    ('номер:', '${passportNumber}'),
    ('кем выдан:', '${passportWho}'),
    ('дата выдачи:', '${passportDate}'),
    ('СНИЛС:', '${SNILS}'),
    ('Телефон:', '${telephone}'),
]
REQ_MINOR_PARENT = [
    ('ФИО (полностью):', '${fullNameParent}'),
    ('Адрес регистрации:', '${passportRegistrationParent}'),
    ('Адрес фактический (почтовый):', '${address}'),
    ('Паспортные данные:', None),
    ('серия:', '${passportSerialParent}'),
    ('номер:', '${passportNumberParent}'),
    ('кем выдан:', '${passportWhoParent}'),
    ('дата выдачи:', '${passportDateParent}'),
    ('код подразделения:', '${passportCodeParent}'),
    ('Телефон:', '${telephoneParent}'),
]
REQ_MINOR_CHILD = [
    ('ФИО (полностью):', '${fullName}'),
    ('Дата рождения:', '${birthDate}'),
    ('Адрес регистрации:', '${passportRegistration}'),
    ('СНИЛС:', '${SNILS}'),
    ('Свидетельство о рождении / паспортные данные:', None),
    ('серия:', '${passportSerial}'),
    ('номер:', '${passportNumber}'),
    ('кем выдано:', '${passportWho}'),
    ('дата выдачи:', '${passportDate}'),
]


def requisites_lines(cell, title, pairs):
    para(cell, title, bold=True, space_after=4)
    for label, value in pairs:
        if value is None:
            para(cell, label, space_before=4, space_after=2)
        else:
            para(cell, '%s %s' % (label, value), size=SMALL_PT + 2, space_after=2)


def build_contract(minor):
    doc = new_document()
    age = 'minor' if minor else 'adult'

    heading(doc, 'Договор №', size=BODY_PT + 2, space_before=0, space_after=2)
    heading(doc, 'безвозмездного оказания услуг', space_before=0, space_after=2)
    heading(doc, 'по социальной интеграции лиц с ограничениями жизнедеятельности',
            space_before=0, space_after=12)

    place = grid(doc, [Cm(8.5), Cm(8.5)])
    cell_margins(place, left=0, right=0)
    row = add_row(place)
    para(row.cells[0], 'г. Москва')
    para(row.cells[1], '«___» ___________ 20___ г.',
         align=WD_ALIGN_PARAGRAPH.RIGHT)
    para(doc, '', space_after=8)

    for kind, text in CONTRACT_BODY[age]:
        if kind == 'h':
            if text.startswith('Терминология'):
                heading(doc, 'Терминология', space_before=12, space_after=0)
                caption(doc, '(термины и определения, используемые в настоящем Договоре)')
            else:
                heading(doc, text, space_before=12)
        elif kind == 'f':
            p = para(doc, text, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=4)
            bottom = OxmlElement('w:bottom')
            bottom.set(qn('w:val'), 'single')
            bottom.set(qn('w:sz'), '4')
            bottom.set(qn('w:space'), '1')
            bottom.set(qn('w:color'), '000000')
            p_pr_add(p, 'w:pBdr').append(bottom)
        elif kind == 'c':
            caption(doc, text)
        else:
            body(doc, text)

    heading(doc, '7. Реквизиты и подписи Сторон', space_before=14, space_after=8)

    table = grid(doc, [Cm(8.5), Cm(8.5)])
    table_borders(table, inside_h=False)
    cell_margins(table, left=142, right=142, top=80, bottom=80)
    row = add_row(table)
    left_cell, right_cell = row.cells

    para(left_cell, 'ИСПОЛНИТЕЛЬ:', bold=True, space_after=4)
    para(left_cell, ORG_FULL, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
         size=SMALL_PT + 2, space_after=4)
    for line in ORG_REQUISITES:
        para(left_cell, line, align=WD_ALIGN_PARAGRAPH.JUSTIFY,
             size=SMALL_PT + 2, space_after=2)
    para(left_cell, 'Заместитель директора', space_before=14, space_after=10)
    para(left_cell, SIGN_LINE, space_after=6)
    para(left_cell, 'м.п.')

    if minor:
        requisites_lines(right_cell, 'ЗАКАЗЧИК:', REQ_MINOR_PARENT)
        para(right_cell, '', space_after=6)
        requisites_lines(right_cell, 'ПОЛУЧАТЕЛЬ УСЛУГ:', REQ_MINOR_CHILD)
        signer = '${fullNameParent}'
    else:
        requisites_lines(right_cell, 'ЗАКАЗЧИК:', REQ_ADULT)
        signer = '${fullName}'
    para(right_cell, '___________________ / %s /' % signer,
         space_before=14, space_after=2)
    caption(right_cell, '(подпись и Ф.И.О. Заказчика)', WD_ALIGN_PARAGRAPH.LEFT)
    return doc


BUILDERS = [
    ('enroll_statement_adult.docx', build_statement, False),
    ('enroll_statement_minor.docx', build_statement, True),
    ('enroll_plan_adult.docx', build_plan, False),
    ('enroll_plan_minor.docx', build_plan, True),
    ('enroll_contract_adult.docx', build_contract, False),
    ('enroll_contract_minor.docx', build_contract, True),
]

if __name__ == '__main__':
    for name, builder, minor in BUILDERS:
        path = os.path.join(OUT, name)
        builder(minor).save(path)
        print('%-30s %7d bytes' % (name, os.path.getsize(path)))
