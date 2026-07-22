# -*- coding: utf-8 -*-
"""
One-time prep: insert ${token} markers into the 4 Word templates at the
fill-in blanks, and copy the 2 Excel templates (which already carry ${tokens})
into test/templates/documents/ with ASCII filenames.

Runtime generation (Node/pizzip) then just does ${token} -> value string
replacement inside the .xml entries of these prepared files.

Each blank is filled by replacing only the underscore run(s) inside the
ORIGINAL paragraph text, so the surrounding label and leading indentation
(the right-aligned addressee block) are preserved verbatim.
"""
import os, re, shutil
from docx import Document
from docx.oxml.ns import qn
from docx.table import Table
from docx.text.paragraph import Paragraph

BASE = r"C:\Users\NIKITA\Desktop\ФАЙЛЫ РАБОТА"
OUT  = r"C:\Users\NIKITA\Desktop\test\templates\documents"
os.makedirs(OUT, exist_ok=True)


def walk(doc):
    """Reproduce the dump numbering: idx increments for each NON-EMPTY
    paragraph and each table, in document order."""
    body = doc.element.body
    idx = 0
    paras, tables = {}, {}
    for child in body.iterchildren():
        if child.tag == qn('w:p'):
            para = Paragraph(child, doc)
            if para.text.strip():
                paras[idx] = para
                idx += 1
        elif child.tag == qn('w:tbl'):
            tables[idx] = Table(child, doc)
            idx += 1
    return paras, tables


def set_para(p, text):
    """Rewrite the paragraph to a single run keeping the first run's font."""
    if not p.runs:
        p.add_run(text)
        return
    p.runs[0].text = text
    for r in list(p.runs[1:]):
        r._element.getparent().remove(r._element)


def set_cell(cell, text):
    set_para(cell.paragraphs[0], text)


def fill_blanks(text, tokens):
    """Replace successive runs of underscores with the given tokens, in order,
    keeping the label and leading indentation. Collapse line breaks to spaces."""
    it = iter(tokens)

    def repl(m):
        try:
            return next(it)
        except StopIteration:
            return m.group(0)

    out = re.sub(r'_+', repl, text)
    out = out.replace('\n', ' ').replace('\r', ' ')
    out = re.sub(r' +([.,;])', r'\1', out)   # drop space before punctuation
    return out


# edit ops: ('fill', [tokens...]) replace blanks in place; ('clear',) empty it
def process(src, dst, para_edits, table_edits):
    doc = Document(os.path.join(BASE, src))
    paras, tables = walk(doc)
    for idx, op in para_edits.items():
        if idx not in paras:
            raise SystemExit(f"[{src}] paragraph idx {idx} not found (have {sorted(paras)})")
        p = paras[idx]
        if op[0] == 'fill':
            set_para(p, fill_blanks(p.text, op[1]))
        elif op[0] == 'clear':
            set_para(p, '')
    for (idx, r, c), text in table_edits.items():
        if idx not in tables:
            raise SystemExit(f"[{src}] table idx {idx} not found (have {sorted(tables)})")
        set_cell(tables[idx].rows[r].cells[c], text)
    doc.save(os.path.join(OUT, dst))
    print("saved", dst)


REL = ('fill', ['${rehRelation}', '${rehRegAddress}'])
CHILD_DOC = ('fill', ['${rehDoc}', ''])   # value line + trailing wrap blank

# ── File 1: PD consent, minor (parent signs) ──────────────────────────────
process(
    "1_SOGLASIE_roditelya_PD_nesovershennoletnego_FORMA_1.docx",
    "pd_consent_minor.docx",
    {
        1: ('fill', ['${repFullName}']),
        3: ('fill', ['${repPassport}']),
        5: ('clear',),
        6: ('fill', ['${repAddress}']),
        7: ('clear',),
        8: ('fill', ['${repPhone}']),
        14: ('fill', ['${rehBirthDate}']),
        16: REL,
        17: CHILD_DOC,
    },
    {
        (11, 0, 1): "${repFullName}",
        (13, 0, 0): "${rehFullName}",
    },
)

# ── File 2: PD consent, 18+ (self) ────────────────────────────────────────
process(
    "2_SOGLASIE_PD_dlya_lits_ot_14_let_FORMA_1.docx",
    "pd_consent_adult.docx",
    {
        1: ('fill', ['${rehFullName}']),
        3: ('fill', ['${rehBirthDate}']),
        4: ('fill', ['${rehPassport}']),
        6: ('clear',),
        7: ('fill', ['${rehRegAddress}']),
        8: ('clear',),
        9: ('fill', ['${rehPhone}']),
    },
    {
        (12, 0, 1): "${rehFullName}",
    },
)

# ── File 3: photo/video consent, 18+ (self) ───────────────────────────────
process(
    "3_SOGLASIE_Reabilitantov_na_foto_video_semku_ot_14_let_FORMA_s_10.docx",
    "photo_consent_adult.docx",
    {
        1: ('fill', ['${rehFullName}']),
        3: ('fill', ['${rehBirthDate}']),
        4: ('fill', ['${rehPassport}']),
        6: ('clear',),
        7: ('fill', ['${rehRegAddress}']),
        8: ('clear',),
        9: ('fill', ['${rehPhone}']),
    },
    {
        (13, 0, 1): "${rehFullName}",
    },
)

# ── File 4: photo/video consent, minor (parent signs) ─────────────────────
process(
    "4_SOGLASIE_Roditelya_reabilitanta_na_foto_video_semku_FORMA_s_10_04.docx",
    "photo_consent_minor.docx",
    {
        1: ('fill', ['${repFullName}']),
        3: ('fill', ['${repPassport}']),
        5: ('clear',),
        6: ('fill', ['${repAddress}']),
        7: ('clear',),
        8: ('fill', ['${repPhone}']),
        15: ('fill', ['${rehBirthDate}']),
        17: REL,
        18: CHILD_DOC,
    },
    {
        (12, 0, 1): "${repFullName}",
        (14, 0, 0): "${rehFullName}",
    },
)

# ── Excel templates already carry ${tokens}; copy verbatim ────────────────
shutil.copy(os.path.join(BASE, "5_Заявление на диагностику 18-.xlsx"),
            os.path.join(OUT, "diagnostics_minor.xlsx"))
print("saved diagnostics_minor.xlsx")
shutil.copy(os.path.join(BASE, "6_Заявление на диагностику 18+.xlsx"),
            os.path.join(OUT, "diagnostics_adult.xlsx"))
print("saved diagnostics_adult.xlsx")

# ── Verify: re-dump the tokenized paragraphs/cells ────────────────────────
print("\n--- verification ---")
for fn in ["pd_consent_minor.docx", "pd_consent_adult.docx",
           "photo_consent_adult.docx", "photo_consent_minor.docx"]:
    d = Document(os.path.join(OUT, fn))
    toks = []
    for p in d.paragraphs:
        if "${" in p.text:
            toks.append(p.text.strip())
    for t in d.tables:
        for row in t.rows:
            for cell in row.cells:
                if "${" in cell.text:
                    toks.append("CELL:" + cell.text.strip())
    print(f"\n{fn}:")
    for t in toks:
        print("   ", t)
print("\ndone")
